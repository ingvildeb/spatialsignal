"""Visual and tabular quality control for within-plane colocalization."""

from __future__ import annotations

import os
from pathlib import Path
import re

import numpy as np
import pandas as pd
from PIL import Image
from skimage.draw import disk, line

from spatialsignal.models import PointCloudDataset
from spatialsignal.pointcloud.colocalization import ColocalizationResult
from spatialsignal.pointcloud._indexing import coordinate_offset


_A_COLOR = np.array([0, 220, 255, 255], dtype=np.uint8)
_B_COLOR = np.array([255, 0, 200, 255], dtype=np.uint8)
_A_UNMATCHED_COLOR = np.array([0, 220, 255, 55], dtype=np.uint8)
_B_UNMATCHED_COLOR = np.array([255, 0, 200, 55], dtype=np.uint8)
_MATCH_LINE_COLOR = np.array([255, 220, 0, 190], dtype=np.uint8)


def write_colocalization_images(
    dataset_a: PointCloudDataset,
    dataset_b: PointCloudDataset,
    result: ColocalizationResult,
    output_dir: Path,
    *,
    output_stem: str,
    marker_radius_px: int = 2,
) -> dict[int, Path]:
    """Write one coordinate-aligned transparent PNG for every declared z plane."""

    if marker_radius_px < 1:
        raise ValueError(f"marker_radius_px must be >= 1, got {marker_radius_px}")
    output_dir.mkdir(parents=True, exist_ok=True)
    for stale_path in output_dir.glob(f"{output_stem}_colocalization_z*.png"):
        stale_path.unlink()
    image_shape = _yx_image_shape(dataset_a)
    offset = coordinate_offset(dataset_a.space.indexing)
    z_values = [int(z) for z in result.plane_summary["z"]]
    z_width = max(4, max((len(str(abs(z))) for z in z_values), default=1))
    image_paths: dict[int, Path] = {}

    for z in z_values:
        a_points = dataset_a.points.loc[dataset_a.points["z"] == z]
        b_points = dataset_b.points.loc[dataset_b.points["z"] == z]
        matches = result.matches.loc[result.matches["z"] == z]
        rgba = np.zeros((*image_shape, 4), dtype=np.uint8)

        a_lookup = a_points.set_index("detection_id", drop=False)
        b_lookup = b_points.set_index("detection_id", drop=False)
        for match in matches.itertuples(index=False):
            a_point = a_lookup.loc[int(match.a_detection_id)]
            b_point = b_lookup.loc[int(match.b_detection_id)]
            ax, ay = _image_xy(a_point, offset)
            bx, by = _image_xy(b_point, offset)
            rr, cc = line(ay, ax, by, bx)
            in_bounds = (
                (rr >= 0)
                & (rr < image_shape[0])
                & (cc >= 0)
                & (cc < image_shape[1])
            )
            rgba[rr[in_bounds], cc[in_bounds]] = _MATCH_LINE_COLOR

        matched_a = set(matches["a_detection_id"].astype(int))
        matched_b = set(matches["b_detection_id"].astype(int))
        _paint_points(
            rgba,
            a_points.loc[~a_points["detection_id"].isin(matched_a)],
            offset=offset,
            radius=marker_radius_px,
            color=_A_UNMATCHED_COLOR,
        )
        _paint_points(
            rgba,
            b_points.loc[~b_points["detection_id"].isin(matched_b)],
            offset=offset,
            radius=marker_radius_px,
            color=_B_UNMATCHED_COLOR,
        )
        _paint_points(
            rgba,
            a_points.loc[a_points["detection_id"].isin(matched_a)],
            offset=offset,
            radius=marker_radius_px,
            color=_A_COLOR,
        )
        _paint_points(
            rgba,
            b_points.loc[b_points["detection_id"].isin(matched_b)],
            offset=offset,
            radius=marker_radius_px,
            color=_B_COLOR,
        )

        sign = "-" if z < 0 else ""
        z_label = f"{sign}{abs(z):0{z_width}d}"
        image_path = output_dir / f"{output_stem}_colocalization_z{z_label}.png"
        Image.fromarray(rgba, mode="RGBA").save(image_path)
        image_paths[z] = image_path

    return image_paths


def write_colocalization_report(
    plane_summary: pd.DataFrame,
    image_paths: dict[int, Path],
    output_path: Path,
    *,
    channel_a_name: str = "a",
    channel_b_name: str = "b",
) -> Path:
    """Write the filterable per-plane XLSX report with links to QC PNGs."""

    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font

    output_path.parent.mkdir(parents=True, exist_ok=True)
    a_name = _excel_column_label(channel_a_name)
    b_name = _excel_column_label(channel_b_name)
    renamed_columns = {
        "n_a_detections": f"n_{a_name}_detections",
        "n_b_detections": f"n_{b_name}_detections",
        "n_unmatched_a": f"n_unmatched_{a_name}",
        "n_unmatched_b": f"n_unmatched_{b_name}",
        "raw_match_rate_a": f"raw_match_rate_{a_name}",
        "raw_match_rate_b": f"raw_match_rate_{b_name}",
        "n_ambiguous_a": f"n_ambiguous_{a_name}",
        "n_ambiguous_b": f"n_ambiguous_{b_name}",
    }
    report = plane_summary.rename(columns=renamed_columns).copy()
    report["png_file"] = [
        os.path.relpath(image_paths[int(z)], output_path.parent).replace("\\", "/")
        for z in report["z"]
    ]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        report.to_excel(writer, sheet_name="planes", index=False)
        worksheet = writer.book["planes"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 30
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(max(len(value) for value in values) + 2, 34)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width

        headers = {cell.value: cell.column for cell in worksheet[1]}
        for row_index in range(2, worksheet.max_row + 1):
            for rate_column in (
                f"raw_match_rate_{a_name}",
                f"raw_match_rate_{b_name}",
            ):
                worksheet.cell(row=row_index, column=headers[rate_column]).number_format = "0.0%"
            png_cell = worksheet.cell(row=row_index, column=headers["png_file"])
            png_cell.hyperlink = png_cell.value
            png_cell.style = "Hyperlink"

        for ambiguity_column in (
            f"n_ambiguous_{a_name}",
            f"n_ambiguous_{b_name}",
        ):
            column_letter = worksheet.cell(
                row=1,
                column=headers[ambiguity_column],
            ).column_letter
            worksheet.conditional_formatting.add(
                f"{column_letter}2:{column_letter}{worksheet.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="FFFFFF",
                    end_type="max",
                    end_color="F4B183",
                ),
            )
    return output_path


def _excel_column_label(channel_name: str) -> str:
    label = re.sub(r"[^A-Za-z0-9]+", "_", channel_name.strip()).strip("_")
    if not label:
        raise ValueError("Channel names must contain at least one alphanumeric character")
    return label


def _paint_points(
    image: np.ndarray,
    points: pd.DataFrame,
    *,
    offset: int,
    radius: int,
    color: np.ndarray,
) -> None:
    for point in points.itertuples(index=False):
        x, y = _image_xy(point, offset)
        rr, cc = disk((y, x), radius=radius, shape=image.shape[:2])
        image[rr, cc] = color


def _image_xy(point: object, offset: int) -> tuple[int, int]:
    x_value = getattr(point, "x_float", getattr(point, "x"))
    y_value = getattr(point, "y_float", getattr(point, "y"))
    x = int(np.floor(float(x_value) - offset + 0.5))
    y = int(np.floor(float(y_value) - offset + 0.5))
    return x, y


def _yx_image_shape(dataset: PointCloudDataset) -> tuple[int, int]:
    shape_by_axis = dict(zip(dataset.space.axis_labels, dataset.space.shape, strict=True))
    return int(shape_by_axis["y"]), int(shape_by_axis["x"])
